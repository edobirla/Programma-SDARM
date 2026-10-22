"""Motore della webapp: serve il pannello (/), l'uscita (/output), il
telefono (/m) e le API.

I dati stanno in una cartella propria (vuota alla prima apertura):
  ~/Library/Application Support/Gestione programma SDARM 2/
    config/settings.json, config/prayers.json
    library/  innari, modelli grafici, lezioni, suoni, immagini…
    cache/    diapositive degli inni salvate come immagini
I file scelti dall'utente vengono COPIATI in library/ e salvati come
percorsi "@/library/…": così l'app non si rompe se l'originale si sposta,
e il backup è una sola cartella.

    python3 webapp/server.py            # http://127.0.0.1:8765
"""
import hashlib
import importlib.util
import json
import mimetypes
import os
import queue
import re
import shutil
import subprocess
import sys
import threading
from http.server import ThreadingHTTPServer, BaseHTTPRequestHandler
from urllib.parse import urlparse, parse_qs, quote

FROZEN = bool(getattr(sys, "frozen", False))      # dentro l'app per Mac
ROOT = sys._MEIPASS if FROZEN else os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = os.environ.get("SDARM_DATA_DIR") or os.path.expanduser(
    "~/Library/Application Support/Gestione programma SDARM 2")
os.environ["SDARM_DATA_DIR"] = DATA          # core.paths legge da qui
os.makedirs(os.path.join(DATA, "library"), exist_ok=True)
sys.path.insert(0, ROOT)
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from core import settings as cfg          # noqa: E402
from core import prayers as prayers_store  # noqa: E402
from core import lesson as lesson_mod      # noqa: E402
from core import bible_ref                 # noqa: E402
from core.bible import BibleLibrary        # noqa: E402
from core.database import load_database    # noqa: E402
import pptx_render                         # noqa: E402
from pptx import Presentation              # noqa: E402

STATIC = os.path.join(ROOT, "webapp", "static")
ASSETS = os.path.join(ROOT, "assets")
PORT = int(os.environ.get("SDARM_PORT") or os.environ.get("PORT") or "8765")
VERSION = "2.0.7"
# identifica questa copia dell'app: l'avvio spegne i motori di versioni precedenti
BUILD = str(int(os.path.getmtime(sys.executable if FROZEN else os.path.abspath(__file__))))
AUDIO_EXT = (".mp3", ".m4a", ".wav", ".ogg", ".aac", ".flac")
SLIDES = os.path.join(DATA, "cache", "slides")

_lock = threading.Lock()
settings = cfg.load()
bible = BibleLibrary()
bible.load_active_from_settings(settings.get("bible_versions", {}))
_hymn_cache = {}          # folder -> [hymn dict]
live = {"v": 0, "state": {"kind": "none"}}
_clients = []             # queue.Queue per SSE client
_jobs = {}                # download da YouTube in corso o appena finiti


# ── percorsi: "@/…" = dentro la cartella dati ────────────────────────────
def R(p):
    """Absolute path for a stored path ("@/library/x" or an absolute one)."""
    if not p:
        return ""
    return os.path.join(DATA, p[2:]) if p.startswith("@/") else p


def rel(p):
    if not p:
        return ""
    p = os.path.abspath(p)
    return "@/" + os.path.relpath(p, DATA) if p.startswith(DATA + os.sep) else p


def save_settings():
    with _lock:
        cfg.save(settings)


# ── importazione nella libreria ──────────────────────────────────────────
def _unique(dest):
    base, ext = os.path.splitext(dest)
    n = 2
    while os.path.exists(dest):
        dest = f"{base} ({n}){ext}"
        n += 1
    return dest


def import_into(kind, src):
    """Copy a file or a folder into library/<kind>/ and return its "@/" path."""
    folder = os.path.join(DATA, "library", kind)
    os.makedirs(folder, exist_ok=True)
    src = src.rstrip("/")
    dest = _unique(os.path.join(folder, os.path.basename(src)))
    if os.path.isdir(src):
        # innari e cartelle musica: solo i file che servono, niente spazzatura
        keep = (".pptx", ".ppt", ".key") + AUDIO_EXT + (".png", ".jpg", ".jpeg")
        shutil.copytree(src, dest, ignore=lambda d, names: [
            n for n in names if n.startswith((".", "~"))
            or not (os.path.isdir(os.path.join(d, n)) or n.lower().endswith(keep))])
    else:
        shutil.copy2(src, dest)
    return rel(dest)


# ── uscite sui proiettori: finestre di Chrome a tutto schermo ───────────────
_outputs = {"n": 0}


def mac_screens():
    """Schermi collegati, il primo è quello con la barra dei menu (il pannello).
    Coordinate in alto a sinistra, come le vuole Chrome."""
    js = ('ObjC.import("AppKit");var s=$.NSScreen.screens,o=[];for(var i=0;i<s.count;i++){var f=s.objectAtIndex(i).frame;'
          'o.push([f.origin.x,f.origin.y,f.size.width,f.size.height])}JSON.stringify(o)')
    try:
        fr = json.loads(subprocess.run(["osascript", "-l", "JavaScript", "-e", js], capture_output=True, text=True, timeout=10).stdout)
    except Exception:
        return []
    main_h = fr[0][3]
    return [{"left": x, "top": main_h - (y + h), "width": w, "height": h} for x, y, w, h in fr]


NATIVE = None          # app per Mac: le uscite sono finestre dell'app (vedi app_main)


def outputs_n():
    """Quante finestre dei proiettori sono aperte (app per Mac: le conta lei)."""
    return NATIVE["count"]() if NATIVE else _outputs["n"]


def open_outputs(screens=None):
    """Una finestra di Chrome in modalità chiosco (schermo intero, niente
    barre) per ogni proiettore, già sullo schermo giusto: nessun clic.
    Con un solo schermo: una finestra normale da spostare dove serve."""
    if NATIVE:
        return NATIVE["open"]()
    close_outputs()
    screens = screens or mac_screens()[1:]
    if not screens:
        prof = os.path.join(DATA, "cache", "chrome", "uscita-0")
        os.makedirs(prof, exist_ok=True)
        subprocess.Popen(["open", "-na", "Google Chrome", "--args", f"--user-data-dir={prof}", "--no-first-run", "--no-default-browser-check",
                          "--autoplay-policy=no-user-gesture-required", "--window-size=960,540", f"--app=http://127.0.0.1:{PORT}/output"])
        return 0
    for i, sc in enumerate(screens):
        prof = os.path.join(DATA, "cache", "chrome", f"uscita-{i}")
        os.makedirs(prof, exist_ok=True)
        subprocess.Popen(["open", "-na", "Google Chrome", "--args", f"--user-data-dir={prof}", "--no-first-run",
                          "--no-default-browser-check", "--autoplay-policy=no-user-gesture-required", "--kiosk",
                          f"--window-position={int(sc['left'])},{int(sc['top'])}", f"--window-size={int(sc['width'])},{int(sc['height'])}",
                          f"--app=http://127.0.0.1:{PORT}/output"])
    return len(screens)


def close_outputs():
    if NATIVE:
        return NATIVE["close"]()
    subprocess.run(["pkill", "-f", "user-data-dir=" + os.path.join(DATA, "cache", "chrome", "uscita-")], check=False)


def search_hymns(q, limit=8, only=""):
    """Tutti gli innari, non solo quello scelto: prima i titoli, poi i testi."""
    q = q.lower().strip()
    title, text = [], []
    for db in ([only] if only else settings.get("databases", [])):
        if not os.path.isdir(R(db)):
            continue
        name = os.path.basename(R(db).rstrip("/"))
        for h in hymns_for(db):
            row = dict(h, db=db, dbname=name)
            if q in (h["number"] + " " + h["title"]).lower():
                title.append(row)
            elif len(q) > 2 and q in " ".join(h["slides"]).lower():
                text.append(row)
    return (title[:limit // 2 + 2] + text)[:limit]


# ── caratteri caricati nell'app (es. Berlin Sans FB copiato da Windows) ────
FONT_EXT = (".ttf", ".otf", ".ttc")


def font_names(path):
    """(famiglia, stile) dalla tabella 'name' del file TTF/OTF."""
    import struct
    with open(path, "rb") as f:
        data = f.read()
    off = 12 if data[:4] != b"ttcf" else struct.unpack(">I", data[12:16])[0] + 12
    n = struct.unpack(">H", data[off - 8:off - 6])[0]
    names = {}
    for i in range(n):
        tag, _, start, _ = struct.unpack(">4sIII", data[off + 16 * i:off + 16 * i + 16])
        if tag != b"name":
            continue
        _, count, strs = struct.unpack(">HHH", data[start:start + 6])
        for j in range(count):
            pid, _, lang, nid, ln, o = struct.unpack(">HHHHHH", data[start + 6 + 12 * j:start + 18 + 12 * j])
            raw = data[start + strs + o:start + strs + o + ln]
            if nid in (1, 2, 16, 17) and (pid == 3 and lang == 0x409 or nid not in names):
                names[nid] = raw.decode("utf-16-be" if pid in (0, 3) else "latin-1", "ignore")
    return names.get(16) or names.get(1, os.path.basename(path)), names.get(17) or names.get(2, "Regular")


def fonts_loaded():
    d = os.path.join(DATA, "library", "fonts")
    out = []
    for f in sorted(os.listdir(d)) if os.path.isdir(d) else []:
        if f.lower().endswith(FONT_EXT):
            try:
                fam, style = font_names(os.path.join(d, f))
            except Exception:
                fam, style = os.path.splitext(f)[0], "Regular"
            out.append({"family": fam, "style": style, "path": rel(os.path.join(d, f))})
    return out


def fonts_css():
    css = []
    for f in fonts_loaded():
        st = f["style"].lower()
        w = 700 if "bold" in st else 600 if "demi" in st or "semi" in st else 300 if "light" in st else 400
        css.append(f"@font-face {{ font-family: '{f['family']}'; src: url('/media?p={quote(f['path'])}'); "
                   f"font-weight: {w}; font-style: {'italic' if 'italic' in st else 'normal'}; }}")
    return "\n".join(css)


def fonts_used():
    """Caratteri usati negli innari (dai primi inni di ognuno)."""
    used = set()
    for db in settings.get("databases", []):
        folder = R(db)
        if not os.path.isdir(folder):
            continue
        for f in sorted(x for x in os.listdir(folder) if x.lower().endswith(".pptx"))[:4]:
            try:
                d = pptx_render.describe(os.path.join(folder, f), SLIDES)
            except Exception:
                continue
            used |= {r["f"] for sl in d["slides"] for sh in sl if "text" in sh for p in sh["text"]["paras"] for r in p["runs"] if r["f"]}
    return sorted(used)


# ── aggiornamenti (versioni 2.x pubblicate su GitHub) ──────────────────────
REPO = "edobirla/Programma-SDARM"
_upd = {"at": 0, "data": None}


def _get(url, timeout=20):
    import ssl
    import urllib.request
    try:
        import certifi
        ctx = ssl.create_default_context(cafile=certifi.where())
    except ImportError:
        ctx = ssl.create_default_context()
    req = urllib.request.Request(url, headers={"User-Agent": "SDARM", "Accept": "application/vnd.github+json"})
    return urllib.request.urlopen(req, timeout=timeout, context=ctx)


def vtuple(v):
    return tuple(int(x) for x in re.findall(r"\d+", v)[:3])


def check_update():
    """Ultima versione 2.x su GitHub (al massimo una richiesta all'ora)."""
    import time as _t
    if _upd["data"] and _t.time() - _upd["at"] < 3600:
        return _upd["data"]
    out = {"current": VERSION, "latest": VERSION, "newer": False, "notes": "", "asset": ""}
    try:
        with _get(f"https://api.github.com/repos/{REPO}/releases?per_page=15") as r:
            rels = json.load(r)
        for rel_ in rels:
            tag = rel_.get("tag_name", "")
            if rel_.get("draft") or rel_.get("prerelease") or not tag.lstrip("v").startswith("2"):
                continue
            asset = next((a["browser_download_url"] for a in rel_.get("assets", []) if a["name"].endswith("-mac.zip")), "")
            out.update(latest=tag.lstrip("v"), notes=rel_.get("body") or "", asset=asset,
                       newer=vtuple(tag) > vtuple(VERSION) and bool(asset))
            break
    except Exception as ex:
        out["error"] = str(ex)
    _upd.update(at=_t.time(), data=out)
    return out


UPD = {"phase": "", "pct": 0.0, "error": ""}     # avanzamento dell'aggiornamento, per la barra nel pannello
APPS = "/Applications/Programma SDARM.app"


def _original_app(app):
    """Dove sta davvero l'app. Aperta da Download (o da uno zip) macOS la fa
    girare da una copia temporanea in sola lettura («AppTranslocation»):
    sostituire quella non serve a niente, va cercata quella vera."""
    if "/AppTranslocation/" not in app:
        return app
    try:
        import ctypes
        cf = ctypes.cdll.LoadLibrary("/System/Library/Frameworks/CoreFoundation.framework/CoreFoundation")
        sec = ctypes.cdll.LoadLibrary("/System/Library/Frameworks/Security.framework/Security")
        cf.CFURLCreateFromFileSystemRepresentation.restype = ctypes.c_void_p
        cf.CFURLCreateFromFileSystemRepresentation.argtypes = [ctypes.c_void_p, ctypes.c_char_p, ctypes.c_long, ctypes.c_bool]
        cf.CFURLGetFileSystemRepresentation.argtypes = [ctypes.c_void_p, ctypes.c_bool, ctypes.c_char_p, ctypes.c_long]
        sec.SecTranslocateCreateOriginalPathForURL.restype = ctypes.c_void_p
        sec.SecTranslocateCreateOriginalPathForURL.argtypes = [ctypes.c_void_p, ctypes.c_void_p]
        raw = app.encode()
        url = cf.CFURLCreateFromFileSystemRepresentation(None, raw, len(raw), True)
        orig = sec.SecTranslocateCreateOriginalPathForURL(url, None)
        buf = ctypes.create_string_buffer(4096)
        if orig and cf.CFURLGetFileSystemRepresentation(orig, True, buf, 4096):
            return buf.value.decode()
    except Exception:
        pass
    return None


def install_update(asset):
    """Scarica la nuova app in sottofondo (la barra nel pannello segue UPD),
    poi uno script la mette in Applicazioni e la riapre; una vecchia copia
    fuori da Applicazioni (es. in Download) va nel Cestino. I dati restano dove sono."""
    app = os.path.abspath(os.path.join(sys.executable, "..", "..", ".."))
    if not FROZEN or not app.endswith(".app"):
        raise RuntimeError("Questa è la versione di sviluppo: aggiorna con  git pull")
    if UPD["phase"] in ("download", "extract", "restart"):
        return
    UPD.update(phase="download", pct=0.0, error="")
    threading.Thread(target=_install, args=(asset, app), daemon=True).start()


def _install(asset, app):
    import tempfile
    try:
        tmp = tempfile.mkdtemp()
        zpath = os.path.join(tmp, "update.zip")
        with _get(asset, timeout=300) as r, open(zpath, "wb") as f:
            total, done = int(r.headers.get("Content-Length") or 0), 0
            while True:
                chunk = r.read(1 << 16)
                if not chunk:
                    break
                f.write(chunk)
                done += len(chunk)
                if total:
                    UPD["pct"] = done / total
        UPD.update(phase="extract", pct=1.0)
        subprocess.run(["ditto", "-x", "-k", zpath, tmp], check=True)
        new = next(os.path.join(tmp, x) for x in os.listdir(tmp) if x.endswith(".app"))
        subprocess.run(["xattr", "-dr", "com.apple.quarantine", new], check=False)
        orig = _original_app(app)
        in_apps = orig and (orig.startswith("/Applications/") or orig.startswith(os.path.expanduser("~/Applications/")))
        target = orig if in_apps else APPS
        if not in_apps and not os.access("/Applications", os.W_OK):
            target = orig or APPS                      # niente permessi su Applicazioni: si resta dov'è
        old_copy = orig if orig and orig != target else ""
        script = os.path.join(tmp, "update.sh")
        with open(script, "w") as f:
            f.write(f"""#!/bin/bash
exec >>"{tmp}/update.log" 2>&1
while kill -0 {os.getpid()} 2>/dev/null; do sleep 0.3; done
if [ -e "{target}" ]; then mv "{target}" "{tmp}/precedente.app" || {{ open "{target}"; exit 1; }}; fi
if ! mv "{new}" "{target}"; then mv "{tmp}/precedente.app" "{target}"; open "{target}"; exit 1; fi
if [ -n "{old_copy}" ] && [ -e "{old_copy}" ]; then mv "{old_copy}" "$HOME/.Trash/Programma SDARM vecchia $(date +%H%M%S).app"; fi
open "{target}"
""")
        UPD.update(phase="restart")
        subprocess.Popen(["/bin/bash", script], start_new_session=True)
        threading.Timer(1.5, lambda: os._exit(0)).start()
    except Exception as ex:
        UPD.update(phase="error", error=str(ex)[:200])


# ── backup, ripristino, vecchia app ─────────────────────────────────────────
def backup_to(folder):
    """Tutti i dati dell'app (senza la cache) in un unico .zip."""
    import zipfile
    from datetime import date
    dest = _unique(os.path.join(folder, f"Backup Programma SDARM {date.today().isoformat()}.zip"))
    with zipfile.ZipFile(dest, "w", zipfile.ZIP_DEFLATED) as z:
        for d, dirs, files in os.walk(DATA):
            dirs[:] = [x for x in dirs if not (d == DATA and x == "cache")]
            for f in files:
                if not f.startswith("."):
                    full = os.path.join(d, f)
                    z.write(full, os.path.relpath(full, DATA))
    return dest


def reload_all():
    settings.clear()
    settings.update(cfg.load())
    bible.load_active_from_settings(settings.get("bible_versions", {}))
    _hymn_cache.clear()
    for k in ("settings", "schedule", "prayers"):
        send_all({"sync": k, "from": ""})


def restore_from(zpath):
    """Il backup sostituisce i dati; quelli di prima restano accanto
    (cartella «… prima del ripristino»), finché non si ripristina di nuovo."""
    import tempfile
    import zipfile
    with zipfile.ZipFile(zpath) as z:
        names = z.namelist()
        if "config/settings.json" not in names or any(n.startswith("/") or ".." in n.split("/") for n in names):
            raise ValueError("Questo file non è un backup del programma")
        tmp = tempfile.mkdtemp(dir=os.path.dirname(DATA))
        z.extractall(tmp)
    old = DATA + " (prima del ripristino)"
    shutil.rmtree(old, ignore_errors=True)
    os.rename(DATA, old)
    os.rename(tmp, DATA)
    os.makedirs(os.path.join(DATA, "library"), exist_ok=True)
    reload_all()


def migrate_from(cfg_dir):
    """Importa dalla vecchia app (cartella «config» con settings.json):
    impostazioni, scaletta, modelli, preghiere; i file che esistono su questo
    Mac (innari, modelli grafici, suoni, immagini) vengono copiati nell'app."""
    with open(os.path.join(cfg_dir, "settings.json"), encoding="utf-8") as f:
        old = json.load(f)
    old_root = os.path.dirname(cfg_dir.rstrip("/"))
    missing = []

    def bring(path, kind):
        if not path:
            return ""
        full = path if os.path.isabs(path) else os.path.join(old_root, path)
        if os.path.exists(full):
            return import_into(kind, full)
        missing.append(os.path.basename(path.replace("\\", "/")) or path)
        return ""

    skip = {"databases", "last_database", "schedule", "hymn_history", "verse_history", "window_geometry", "lesson_file",
            "lesson_photos", "youtube_folder", "left_panel_width", "right_panel_width", "inni_preview_width", "tutorial_seen",
            "projection_screen", "bible_versions"}
    files = {"bell_sound": "suoni", "timer_warning_sound": "suoni", "offering_image": "immagini", "background_image": "immagini",
             "verse_bg_image": "immagini", "timer_bg_image": "immagini", "background_music_folder": "musica"}
    for k, v in old.items():
        if k in skip:
            continue
        settings[k] = bring(v, files[k]) if k in files else v
    for prof in (settings.get("lesson_graphic_profiles") or {}).values():
        if prof.get("template"):
            prof["template"] = bring(prof["template"], "modelli")
    dbmap = {}
    for d in old.get("databases", []):
        new = bring(d, "innari")
        if new:
            dbmap[os.path.abspath(d if os.path.isabs(d) else os.path.join(old_root, d))] = new
            if new not in settings.setdefault("databases", []):
                settings["databases"].append(new)
    if dbmap:
        settings["last_database"] = dbmap.get(os.path.abspath(old.get("last_database", "")), next(iter(dbmap.values())))

    def fix_slots(secs):
        for sec in secs:
            for sl in sec.get("slots", []):
                h = (sl.get("data") or {}).get("hymn")
                if not h:
                    continue
                for key in ("path", "audio_path"):
                    p = h.get(key) or ""
                    full = os.path.abspath(p if os.path.isabs(p) else os.path.join(old_root, p)) if p else ""
                    for o, n in dbmap.items():
                        if full.startswith(o + os.sep):
                            h[key] = n + "/" + os.path.relpath(full, o)
        return secs
    if old.get("schedule"):
        settings["schedule"] = fix_slots(old["schedule"])
    tdir = os.path.join(cfg_dir, "templates")
    if os.path.isdir(tdir):
        tpls = [t for t in settings.get("schedule_templates", [])]
        for f in sorted(os.listdir(tdir)):
            if f.endswith(".json"):
                with open(os.path.join(tdir, f), encoding="utf-8") as fh:
                    secs = fix_slots(json.load(fh))
                tpls = [t for t in tpls if t["name"] != f[:-5]] + [{"name": f[:-5], "secs": secs}]
        settings["schedule_templates"] = tpls
    if os.path.isfile(os.path.join(cfg_dir, "prayers.json")):
        shutil.copy2(os.path.join(cfg_dir, "prayers.json"), prayers_store.PRAYERS_PATH)
    save_settings()
    reload_all()
    return {"ok": True, "dbs": len(dbmap), "missing": missing}


def pick(kind, prompt, types=None, multiple=False):
    """Native macOS file/folder dialog, opened by the engine on this Mac."""
    prompt = prompt.replace('"', "'")
    if kind == "folder":
        cmd = f'POSIX path of (choose folder with prompt "{prompt}")'
    else:
        of = (" of type {" + ",".join(f'"{t}"' for t in types) + "}") if types else ""
        if multiple:
            cmd = (f'set fs to (choose file with prompt "{prompt}"{of} with multiple selections allowed)\n'
                   'set out to ""\nrepeat with f in fs\nset out to out & POSIX path of f & linefeed\nend repeat\nout')
        else:
            cmd = f'POSIX path of (choose file with prompt "{prompt}"{of})'
    script = f'tell application "System Events"\nactivate\n{cmd}\nend tell'
    r = subprocess.run(["osascript", "-e", script], capture_output=True, text=True)
    if r.returncode != 0:      # annullato
        return []
    return [line for line in r.stdout.strip().splitlines() if line]


# ── inni ─────────────────────────────────────────────────────────────────
def audio_by_number(folder):
    """{9: path} for every audio file whose name starts with a number — so
    "9 Lavami.mp3", "009-lavami.MP3" and "009 - Lavami, Lavami.mp3" all match
    hymn 009 (the desktop app required the exact pptx file name)."""
    out = {}
    for f in os.listdir(folder):
        head = f.split()[0].split("-")[0].split("_")[0] if f.strip() else ""
        if f.lower().endswith(AUDIO_EXT) and head.isdigit():
            out.setdefault(int(head), os.path.join(folder, f))
    return out


def hymns_for(stored):
    folder = R(stored)
    if folder not in _hymn_cache:
        audio = audio_by_number(folder)
        out = []
        for h in load_database(folder):
            a = (h.audio_path if h.audio_path and os.path.isfile(h.audio_path) else None) or (audio.get(int(h.number)) if h.number.isdigit() else None)
            out.append({"number": h.number, "title": h.title, "path": rel(h.file_path),
                        "audio": rel(a) if a else "", "slides": h.slides_text})
        _hymn_cache[folder] = out
    return _hymn_cache[folder]


# ── lezione ──────────────────────────────────────────────────────────────
def lesson_payload():
    path = R(settings.get("lesson_file", ""))
    if not path or not os.path.isfile(path):
        return {"ok": False, "reason": "missing", "file": settings.get("lesson_file", "")}
    try:
        ls = lesson_mod.LessonSet.load(path)
    except Exception as ex:  # file rotto: il pannello lo dice, non si blocca
        return {"ok": False, "reason": str(ex), "file": settings.get("lesson_file", "")}
    idx = settings.get("lesson_override_index", -1)
    les = ls.get(idx) if isinstance(idx, int) and idx >= 0 else ls.find_for()
    out = {"ok": True, "titles": ls.titles(), "index": idx, "lesson": None}
    if not les:
        return out
    lang = settings.get("bible_primary_lang") or (bible.languages() or ["it"])[0]
    books = bible.books(lang)
    domande = lesson_mod.flatten_domande(les)
    for d in domande:
        refs = bible_ref.parse_references(d["verses"], books) if books else []
        for r in refs:
            r["texts"] = [bible.get_verse(lang, r["book"], r["chapter"], v) or "" for v in r["verses"]] \
                if r["resolved"] else []
        d["refs"] = refs
        dd = lesson_mod.day_date(les, d)
        d["date"] = dd.isoformat() if dd else ""
    key = les.get("data_sabato") or les.get("date") or ""
    intro = settings.get("lesson_intro", {}).get(key, {})
    out["lesson"] = {
        "key": key, "number": les.get("numero_lezione"), "title": les.get("title") or les.get("titolo") or "Lezione",
        "trimestre": les.get("trimestre", ""), "sabato": les.get("data_sabato", ""),
        "verse": intro.get("verse") or les.get("versetto_chiave") or les.get("memory_verse") or "",
        "verse_ref": intro.get("ref") or les.get("versetto_chiave_riferimento") or "",
        "note": intro.get("note") or les.get("nota_introduttiva") or "",
        "photo": settings.get("lesson_photos", {}).get(key, ""),
        # prompt per l'immagine di copertina (1920x340), se il file JSON li contiene
        "prompt": les.get("prompt_immagine", ""),
        "domande": domande}
    return out


def broadcast(state):
    with _lock:
        live["v"] += 1
        live["state"] = state
        msg = json.dumps(live)
        for q in list(_clients):
            q.put(msg)
    if BRIDGE:
        BRIDGE.push(msg)


def send_all(obj):
    """Messaggi che non sono lo stato in onda: comandi al Mac, stato audio
    per il telefono, avvisi di dati cambiati."""
    msg = json.dumps(obj)
    for q in list(_clients):
        q.put(msg)
    if BRIDGE and "view" not in obj:          # il puntatore non serve al telefono (e sarebbero troppi messaggi)
        BRIDGE.push(msg)


BRIDGE = None     # ponte verso il telefono con i dati cellulare (cloud_bridge.py)


def phone_ids():
    """Stanza (codice lungo) e codice breve del telefono, creati la prima volta."""
    import secrets
    if not settings.get("phone_room"):
        settings["phone_room"] = os.urandom(16).hex()
        save_settings()
    if not settings.get("phone_code"):
        settings["phone_code"] = "".join(secrets.choice("ABCDEFGHJKLMNPQRSTUVWXYZ23456789") for _ in range(8))
        save_settings()
    if BRIDGE and BRIDGE.room != settings["phone_room"]:       # codice cambiato: il ponte segue
        BRIDGE.room, BRIDGE.code, BRIDGE.exp = settings["phone_room"], settings["phone_code"], 0
    return settings["phone_room"], settings["phone_code"]


# ── telefono: stessa rete Wi-Fi, protetto da un PIN ──────────────────────────
import hmac   # noqa: E402
import time   # noqa: E402
_fails = {"n": 0, "until": 0.0}
CLOUD_URL = "https://gestionale-programmi.web.app"     # pagina del telefono pubblicata su Firebase Hosting
PHONE_BLOCKED = {"/api/library/import", "/api/pick", "/api/library/remove", "/api/hymn/export", "/api/backup/export", "/api/backup/import", "/api/migrate", "/api/update/install"}


def phone_token():
    pin = str(settings.get("phone_pin") or "")
    if not pin:
        return ""
    if not settings.get("phone_secret"):
        settings["phone_secret"] = os.urandom(16).hex()
        save_settings()
    return hmac.new(settings["phone_secret"].encode(), pin.encode(), "sha256").hexdigest()


def lan_address():
    import socket
    try:
        with socket.socket(socket.AF_INET, socket.SOCK_DGRAM) as so:
            so.connect(("10.255.255.255", 1))
            return so.getsockname()[0]
    except OSError:
        return ""


class Handler(BaseHTTPRequestHandler):
    server_version = "SDARM/" + VERSION

    def log_message(self, *a):
        pass

    # ── helpers ────────────────────────────────────────────────────────────
    def _json(self, obj, code=200):
        body = json.dumps(obj, ensure_ascii=False).encode("utf-8")
        self.send_response(code)
        self.send_header("X-SDARM-Build", BUILD)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        self.wfile.write(body)

    def _body(self):
        n = int(self.headers.get("Content-Length") or 0)
        return json.loads(self.rfile.read(n) or b"null")

    def _file(self, path, ranged=False):
        if not path or not os.path.isfile(path):
            return self.send_error(404)
        size = os.path.getsize(path)
        ctype = mimetypes.guess_type(path)[0] or "application/octet-stream"
        start, end = 0, size - 1
        rng = self.headers.get("Range") if ranged else None
        if rng and rng.startswith("bytes="):
            a, _, b = rng[6:].partition("-")
            start = int(a) if a else max(0, size - int(b))
            end = int(b) if a and b else size - 1
            self.send_response(206)
            self.send_header("Content-Range", f"bytes {start}-{end}/{size}")
        else:
            self.send_response(200)
        self.send_header("Content-Type", ctype)
        self.send_header("Accept-Ranges", "bytes")
        self.send_header("Content-Length", str(end - start + 1))
        self.send_header("Cache-Control", "no-store" if path.startswith(STATIC) else "max-age=600")
        self.end_headers()
        with open(path, "rb") as f:
            f.seek(start)
            left = end - start + 1
            while left > 0:
                chunk = f.read(min(65536, left))
                if not chunk:
                    break
                try:
                    self.wfile.write(chunk)
                except (BrokenPipeError, ConnectionResetError):
                    return
                left -= len(chunk)

    def _bootstrap(self):
        dbs = settings.get("databases", [])
        return {
            "version": VERSION, "data": DATA,
            "settings": {k: v for k, v in settings.items() if k not in ("schedule", "hymn_history", "verse_history", "phone_secret")
                         and (self._local() or k not in ("phone_pin", "phone_room", "phone_code"))},
            "schedule": settings.get("schedule", []),
            "prayers": prayers_store.load(),
            "dbs": [{"path": d, "name": os.path.basename(R(d).rstrip("/")) or d, "exists": os.path.isdir(R(d))} for d in dbs],
            "langs": [{"code": c, "label": bible.language_label(c)} for c in bible.languages()],
            "live": live,
        }

    def _loopback(self):
        return self.client_address[0] in ("127.0.0.1", "::1", "::ffff:127.0.0.1")

    def _local(self):
        # le richieste del telefono via Firebase passano dal pannello del Mac:
        # arrivano da qui ma valgono come telefono (stessi limiti)
        return self._loopback() and not self.headers.get("X-Via-Cloud")

    def _authed(self):
        """Mac: sempre. Telefono in Wi-Fi: con il PIN giusto (cookie).
        Telefono via Firebase: il codice della stanza l'ha già controllato Firebase."""
        if self._loopback():
            return True
        tok = phone_token()
        ck = dict(c.strip().split("=", 1) for c in (self.headers.get("Cookie") or "").split(";") if "=" in c)
        return bool(tok) and hmac.compare_digest(ck.get("sdarm", ""), tok)

    def _gate(self, path):
        if self._authed():
            if not self._local() and path in PHONE_BLOCKED:
                self._json({"ok": False, "error": "Solo dal Mac"}, 403)
                return False
            return True
        if path in ("/m", "/api/phone/login", "/fonts.css") or path.startswith(("/static/", "/assets/")):
            return True
        self._json({"ok": False, "error": "PIN"}, 401)
        return False

    def _sync(self, what):
        send_all({"sync": what, "from": self.headers.get("X-Client", "")})

    # ── routes ─────────────────────────────────────────────────────────────
    def do_GET(self):
        u = urlparse(self.path)
        q = {k: v[0] for k, v in parse_qs(u.query).items()}
        p = u.path
        if not self._gate(p):
            return
        if p == "/":
            return self._file(os.path.join(STATIC, "index.html"))
        if p == "/m":
            return self._file(os.path.join(STATIC, "index.html" if self._authed() else "pin.html"))
        if p == "/api/tempo":
            # base a velocità diversa, preparata da ffmpeg (suona naturale,
            # stessa intonazione); una volta pronta resta in cache
            src, r = R(q.get("p", "")), round(float(q.get("r", "1")), 2)
            if not os.path.isfile(src) or not 0.5 <= r <= 1.5:
                return self._json({"ok": False}, 400)
            st = os.stat(src)
            out = os.path.join(DATA, "cache", "tempo", hashlib.md5(f"{src}|{st.st_mtime}".encode()).hexdigest()[:16] + f"_{int(r * 100)}.m4a")
            if not os.path.isfile(out):
                import imageio_ffmpeg
                os.makedirs(os.path.dirname(out), exist_ok=True)
                tmp = out + ".tmp.m4a"
                subprocess.run([imageio_ffmpeg.get_ffmpeg_exe(), "-y", "-v", "error", "-i", src, "-vn", "-filter:a", f"atempo={r}",
                                "-c:a", "aac", "-b:a", "192k", tmp], check=True, timeout=300)
                os.replace(tmp, out)
            return self._json({"ok": True, "path": rel(out)})
        if p == "/api/search/hymns":
            return self._json(search_hymns(q.get("q", ""), int(q.get("limit", 8)), q.get("db", "")))
        if p == "/fonts.css":
            body = fonts_css().encode()
            self.send_response(200)
            self.send_header("Content-Type", "text/css; charset=utf-8")
            self.send_header("Content-Length", str(len(body)))
            self.send_header("Cache-Control", "no-store")
            self.end_headers()
            return self.wfile.write(body)
        if p == "/api/fonts":
            return self._json({"loaded": fonts_loaded(), "used": fonts_used()})
        if p == "/api/update":
            return self._json(check_update())
        if p == "/api/update/status":
            return self._json(UPD)
        if p == "/api/phone/qr":
            import io
            import qrcode
            import qrcode.image.svg
            buf = io.BytesIO()
            qrcode.make(q.get("u", ""), image_factory=qrcode.image.svg.SvgPathFillImage, box_size=12, border=2).save(buf)
            body = buf.getvalue()
            self.send_response(200)
            self.send_header("Content-Type", "image/svg+xml")
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            return self.wfile.write(body)
        if p == "/api/phone":
            ip = lan_address()
            room, _code = phone_ids()
            return self._json({"url": f"http://{ip}:{PORT}/m" if ip else "", "pin": bool(settings.get("phone_pin")),
                               "room": room if self._local() else "", "code": settings["phone_code"] if self._local() else "",
                               "cloud": f"{CLOUD_URL}/m?r={room}"})
        if p == "/output":
            return self._file(os.path.join(STATIC, "output.html"))
        if p.startswith("/static/"):
            f = os.path.normpath(os.path.join(STATIC, p[len("/static/"):]))
            return self._file(f) if f.startswith(STATIC) else self.send_error(403)
        if p.startswith("/assets/"):
            f = os.path.normpath(os.path.join(ASSETS, p[len("/assets/"):]))
            return self._file(f) if f.startswith(ASSETS) else self.send_error(403)
        if p == "/media":
            # il Mac può aprire qualsiasi file; il telefono solo i dati dell'app
            f = os.path.realpath(R(q.get("p", "")))
            if not self._local() and not f.startswith(os.path.realpath(DATA) + os.sep):
                return self.send_error(403)
            return self._file(f, ranged=True)
        if p == "/api/bootstrap":
            return self._json(self._bootstrap())
        if p == "/api/hymns":
            folder = q.get("db") or settings.get("last_database", "")
            if not folder or not os.path.isdir(R(folder)):
                return self._json({"ok": False, "hymns": []})
            hy = hymns_for(folder)
            if q.get("lite"):            # telefono: solo numero, titolo e base (il testo si cerca sul Mac)
                hy = [{k: h[k] for k in ("number", "title", "path", "audio")} for h in hy]
            return self._json({"ok": True, "hymns": hy})
        if p == "/api/bible/books":
            return self._json(bible.books(q.get("lang", "it")))
        if p == "/api/bible/chapters":
            return self._json(bible.chapters(q.get("lang", "it"), int(q["book"])))
        if p == "/api/bible/verses":
            return self._json(bible.verses(q.get("lang", "it"), int(q["book"]), int(q["ch"])))
        if p == "/api/bible/verse":
            book, ch, v = int(q["book"]), int(q["ch"]), int(q["v"])
            return self._json({c: bible.get_verse(c, book, ch, v) or "" for c in q.get("langs", "it").split(",")})
        if p == "/api/bible/search":
            lang = q.get("lang", "it")
            res = bible.search(lang, q.get("q", ""), limit=int(q.get("limit", 40)))
            return self._json([{"book": b, "name": bible.book_name(lang, b), "ch": c, "v": v, "text": t}
                               for b, c, v, t in res])
        if p == "/api/slides":
            try:
                d = pptx_render.describe(R(q.get("p", "")), SLIDES)
                if d.get("pdf"):                    # l'uscita chiede le pagine con i testi modificati tolti
                    for sl in d["slides"]:
                        sl[0]["src"] = q.get("p", "")
                return self._json(d)
            except Exception as ex:      # file rotto: l'uscita usa il testo
                return self._json({"error": str(ex)}, 422)
        if p == "/api/pdfcut":
            # pagina di un PDF senza le scritte dei riquadri modificati (cache)
            src = os.path.realpath(R(q.get("p", "")))
            if not self._local() and not src.startswith(os.path.realpath(DATA) + os.sep):
                return self.send_error(403)
            n = int(q.get("n", "0"))
            cut = sorted({int(x) for x in q.get("cut", "").split(",") if x.isdigit()})
            shapes = pptx_render.describe(src, SLIDES)["slides"][n]
            key = shapes[0]["img"].split("/")[0]
            out = os.path.join(SLIDES, key, f"cut{n:04d}-{'-'.join(map(str, cut))}.png")
            if not os.path.isfile(out):
                pptx_render.pdf_cut(src, n, [shapes[i]["bb"] for i in cut if i < len(shapes) and "bb" in shapes[i]], out)
            return self._file(out)
        if p == "/api/pdf":
            # pagine del PDF come immagini, una volta sola (cache)
            import fitz
            src = R(q.get("p", ""))
            st = os.stat(src)
            key = hashlib.md5(f"{src}|{st.st_mtime}|{st.st_size}".encode()).hexdigest()[:20]
            out = os.path.join(DATA, "cache", "pdf", key)
            if not os.path.isdir(out):
                os.makedirs(out + ".tmp", exist_ok=True)
                with fitz.open(src) as doc:
                    for i, page in enumerate(doc):
                        z = 1920 / max(page.rect.width, 1)
                        page.get_pixmap(matrix=fitz.Matrix(z, z)).save(os.path.join(out + ".tmp", f"{i:04d}.png"))
                os.replace(out + ".tmp", out)
            return self._json({"pages": [f"@/cache/pdf/{key}/{f}" for f in sorted(os.listdir(out))]})
        if p == "/api/storage":
            def size(d):
                n = 0
                for root, _, fs in os.walk(d):
                    for f in fs:
                        try:
                            n += os.path.getsize(os.path.join(root, f))
                        except OSError:
                            pass
                return n
            parts = [("Libreria (audio, video, immagini)", size(os.path.join(DATA, "library"))),
                     ("Bibbie", size(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "bibles"))),
                     ("Cache e anteprime", size(os.path.join(DATA, "cache")))]
            parts.append(("Altri dati dell'app", max(0, size(DATA) - parts[0][1] - parts[2][1])))
            dbs = [(os.path.basename(d.rstrip("/")) or d, size(R(d))) for d in (settings.get("databases") or []) if isinstance(d, str) and os.path.isdir(R(d))]
            du = shutil.disk_usage(DATA)
            web = size(STATIC) + size(os.path.join(os.path.dirname(STATIC), "..", "assets"))
            fb = None
            if BRIDGE:
                try:
                    fb = len(json.dumps(BRIDGE.db("GET", "") or ""))
                except Exception:
                    pass
            return self._json({"app": parts, "hymnals": dbs, "disk": {"free": du.free, "total": du.total}, "web": web, "fb": fb})
        if p == "/api/hymnal/template":
            folder = q.get("db", "")
            tpl = (settings.get("hymnal_templates") or {}).get(folder) or pptx_render.derive_template(R(folder), SLIDES)
            if tpl.get("bg") and not tpl["bg"].startswith("@/"):
                tpl = dict(tpl, bg="@/cache/slides/" + tpl["bg"])
            return self._json(tpl)
        if p == "/api/youtube/jobs":
            import time as _t
            return self._json([j for j in _jobs.values() if j["status"] == "run" or _t.time() - j["at"] < 3600])
        if p == "/api/library":
            # libreria media: video e audio (scaricati da YouTube o caricati)
            out = []
            for kind in ("video", "audio"):
                d = os.path.join(DATA, "library", kind)
                for f in sorted(os.listdir(d)) if os.path.isdir(d) else []:
                    full = os.path.join(d, f)
                    if os.path.isfile(full) and not f.startswith("."):
                        st = os.stat(full)
                        out.append({"kind": kind, "name": os.path.splitext(f)[0], "path": rel(full), "size": st.st_size, "at": st.st_mtime})
            out.sort(key=lambda x: -x["at"])
            return self._json(out)
        if p == "/api/audiofiles":
            folder = R(q.get("folder", ""))
            files = sorted(f for f in os.listdir(folder) if f.lower().endswith(AUDIO_EXT)) if os.path.isdir(folder) else []
            return self._json([rel(os.path.join(folder, f)) for f in files])
        if p == "/api/lesson":
            return self._json(lesson_payload())
        if p == "/api/live":
            return self._json(dict(live, outputs=outputs_n()))
        if p == "/api/events":
            self.send_response(200)
            self.send_header("Content-Type", "text/event-stream")
            self.send_header("Cache-Control", "no-store")
            self.end_headers()
            qu = queue.Queue()
            _clients.append(qu)
            is_out = q.get("role") == "output"
            if is_out:
                _outputs["n"] += 1
                send_all({"outputs": outputs_n()})
            try:
                self.wfile.write(f"data: {json.dumps(live)}\n\n".encode())
                self.wfile.flush()
                while True:
                    try:
                        msg = qu.get(timeout=15)
                        self.wfile.write(f"data: {msg}\n\n".encode())
                    except queue.Empty:
                        self.wfile.write(b": ping\n\n")
                    self.wfile.flush()
            except (BrokenPipeError, ConnectionResetError, OSError):
                pass
            finally:
                _clients.remove(qu)
                if is_out:
                    _outputs["n"] -= 1
                    send_all({"outputs": outputs_n()})
            return
        self.send_error(404)

    def do_PUT(self):
        p = urlparse(self.path).path
        if not self._gate(p):
            return
        data = self._body()
        if p == "/api/schedule":
            settings["schedule"] = data
            save_settings()
            self._sync("schedule")
            return self._json({"ok": True})
        if p == "/api/prayers":
            prayers_store.save(data)
            self._sync("prayers")
            return self._json({"ok": True})
        if p == "/api/settings":
            for k, v in (data or {}).items():
                if k not in ("schedule", "phone_secret") and (self._local() or k not in ("phone_pin", "phone_room", "phone_code")):
                    settings[k] = v
            save_settings()
            self._sync("settings")
            if "databases" in data:
                _hymn_cache.clear()
                # gli inni si leggono subito, in sottofondo: la prima apertura è già pronta
                threading.Thread(target=lambda: [hymns_for(d) for d in data["databases"] if os.path.isdir(R(d))], daemon=True).start()
            return self._json({"ok": True})
        self.send_error(404)

    def do_POST(self):
        p = urlparse(self.path).path
        if not self._gate(p):
            return
        if p == "/api/phone/login":
            if time.time() < _fails["until"]:
                return self._json({"ok": False, "error": "Troppi tentativi: riprova tra qualche minuto"}, 429)
            tok = phone_token()
            if tok and hmac.compare_digest(str(self._body().get("pin", "")), str(settings.get("phone_pin"))):
                _fails["n"] = 0
                body = b'{"ok": true}'
                self.send_response(200)
                self.send_header("Content-Type", "application/json")
                self.send_header("Content-Length", str(len(body)))
                self.send_header("Set-Cookie", f"sdarm={tok}; Path=/; Max-Age=31536000; HttpOnly; SameSite=Strict")
                self.end_headers()
                return self.wfile.write(body)
            _fails["n"] += 1
            if _fails["n"] >= 8:
                _fails["n"], _fails["until"] = 0, time.time() + 300
            time.sleep(1)
            return self._json({"ok": False, "error": "PIN sbagliato" if tok else "Imposta prima il PIN sul Mac"}, 401)
        if p == "/api/outputs/open":
            if not NATIVE and not os.path.isdir("/Applications/Google Chrome.app"):
                return self._json({"ok": False, "error": "Serve Google Chrome"})
            n = open_outputs(self._body().get("screens"))
            if NATIVE:
                send_all({"outputs": outputs_n()})
            return self._json({"ok": True, "screens": n})
        if p == "/api/outputs/close":
            close_outputs()
            if NATIVE:
                send_all({"outputs": 0})
            return self._json({"ok": True})
        if p == "/api/upload":
            # file trascinati dal Finder nella scaletta: entrano nella libreria
            kind = parse_qs(urlparse(self.path).query).get("kind", ["documenti"])[0]
            name = os.path.basename(parse_qs(urlparse(self.path).query).get("name", ["file"])[0]) or "file"
            if kind not in ("immagini", "video", "audio", "presentazioni", "documenti"):
                return self._json({"ok": False}, 400)
            folder = os.path.join(DATA, "library", kind)
            os.makedirs(folder, exist_ok=True)
            dest = _unique(os.path.join(folder, name))
            left = int(self.headers.get("Content-Length") or 0)
            with open(dest, "wb") as f:
                while left > 0:
                    chunk = self.rfile.read(min(1 << 20, left))
                    if not chunk:
                        break
                    f.write(chunk)
                    left -= len(chunk)
            return self._json({"ok": True, "path": rel(dest)})
        if p == "/api/cmd":
            # telefono -> Mac (comandi) e Mac -> telefono (stato audio e timer)
            send_all(self._body())
            return self._json({"ok": True})
        if p == "/api/live":
            broadcast(self._body())
            return self._json({"ok": True, "v": live["v"]})
        if p == "/api/pick":
            # finestra del Mac per scegliere; poi il file entra nella libreria
            b = self._body()
            paths = pick(b.get("kind", "file"), b.get("prompt", "Scegli"), b.get("types"), b.get("multiple", False))
            if b.get("library"):
                paths = [import_into(b["library"], x) for x in paths]
            return self._json({"paths": paths})
        if p == "/api/hymn/import":
            # PowerPoint / Word / testo -> nuovo inno nello stile dell'innario
            b = self._body()
            folder = R(b["db"])
            tpl = (settings.get("hymnal_templates") or {}).get(b["db"]) or pptx_render.derive_template(folder, SLIDES)
            bg = R(tpl["bg"]) if tpl.get("bg", "").startswith("@/") else (os.path.join(SLIDES, tpl["bg"]) if tpl.get("bg") else "")
            stanzas = b.get("stanzas") or pptx_render.read_blocks(R(b["src"]))
            title = b.get("title") or (stanzas[0].split("\n")[0] if stanzas else "Inno")
            if b.get("src", "").lower().endswith(".pptx") and stanzas and stanzas[0].strip() == title:
                stanzas = stanzas[1:]
            name = f"{str(b.get('number', '')).zfill(3)} - {title}.pptx" if b.get("number") else f"{title}.pptx"
            dest = os.path.join(folder, name.replace("/", "-"))
            pptx_render.build_pptx(dest, title, stanzas, tpl, bg)
            _hymn_cache.pop(folder, None)
            return self._json({"ok": True, "path": rel(dest)})
        if p == "/api/youtube":
            # scarica il video nella libreria (serve yt-dlp sul Mac)
            b = self._body()
            url, audio_only = b.get("url", "").strip(), b.get("kind") == "audio"
            if not url.startswith(("http://", "https://")):
                return self._json({"ok": False, "error": "Indirizzo non valido"})
            dest = os.path.join(DATA, "library", "audio" if audio_only else "video")
            os.makedirs(dest, exist_ok=True)
            try:
                import yt_dlp
            except ImportError:
                return self._json({"ok": False, "error": "Manca yt-dlp: installalo con  pip3 install yt-dlp"})
            opts = {"noplaylist": True, "quiet": True, "no_warnings": True, "outtmpl": os.path.join(dest, "%(title).80s.%(ext)s"),
                    "format": "b[ext=mp4]/b"}
            try:                          # ffmpeg unisce video e audio in un mp4 H.264
                import imageio_ffmpeg
                opts.update(ffmpeg_location=imageio_ffmpeg.get_ffmpeg_exe(), merge_output_format="mp4",
                            format="bv*[vcodec^=avc1]+ba[ext=m4a]/b[ext=mp4]/bv*+ba/b")
            except ImportError:
                pass
            if audio_only:                # solo il suono, in m4a (si apre ovunque)
                opts.update(format="ba[ext=m4a]/ba", merge_output_format=None)
            # in sottofondo: il pannello segue l'avanzamento da /api/youtube/jobs
            import time as _t
            job = {"id": os.urandom(4).hex(), "kind": "audio" if audio_only else "video", "name": "", "pct": 0.0, "status": "run",
                   "error": "", "path": "", "slot": b.get("slot") or "", "at": _t.time()}
            _jobs[job["id"]] = job

            def hook(d):
                job["name"] = job["name"] or (d.get("info_dict") or {}).get("title", "")
                if d.get("status") == "downloading":
                    tot = d.get("total_bytes") or d.get("total_bytes_estimate") or 0
                    job["pct"] = min(0.99, (d.get("downloaded_bytes") or 0) / tot) if tot else job["pct"]
                elif d.get("status") == "finished":
                    job["pct"] = 0.99

            def work():
                try:
                    with yt_dlp.YoutubeDL(dict(opts, progress_hooks=[hook])) as y:
                        info = y.extract_info(url, download=True)
                    path = (info.get("requested_downloads") or [{}])[0].get("filepath", "")
                    if not os.path.isfile(path):
                        raise RuntimeError("Download non riuscito")
                    job.update(status="done", pct=1.0, path=rel(path), name=os.path.splitext(os.path.basename(path))[0])
                except Exception as ex:
                    job.update(status="error", error=str(ex).replace("ERROR: ", "")[:200])
            threading.Thread(target=work, daemon=True).start()
            return self._json({"ok": True, "job": job["id"]})
        if p == "/api/hymn/delete":
            # l'inno va nel Cestino (si può recuperare); sparisce anche dalle modifiche salvate
            b = self._body()
            src = R(b.get("path", ""))
            if not src.lower().endswith(".pptx") or not os.path.isfile(src):
                return self._json({"ok": False}, 400)
            try:   # Cestino di Finder: si può recuperare con «Rimetti a posto»
                subprocess.run(["osascript", "-e", 'tell application "Finder" to delete POSIX file ' + json.dumps(src)],
                               check=True, capture_output=True, timeout=20)
            except Exception:
                return self._json({"ok": False, "error": "Il Mac non mi ha lasciato spostarlo nel Cestino"})
            for k in ("hymn_edits", "hymn_bases"):
                if isinstance(settings.get(k), dict):
                    settings[k].pop(b["path"], None)
            save_settings()
            _hymn_cache.clear()
            return self._json({"ok": True})
        if p == "/api/hymn/slide-delete":
            # toglie una diapositiva dal file dell'inno; le modifiche di testo salvate seguono lo scorrimento
            b = self._body()
            src, i = R(b.get("path", "")), int(b.get("index", -1))
            prs = Presentation(src)
            if not src.lower().endswith(".pptx") or not 0 <= i < len(prs.slides) or len(prs.slides) < 2:
                return self._json({"ok": False}, 400)
            ids = prs.slides._sldIdLst
            prs.part.drop_rel(ids[i].rId)
            ids.remove(ids[i])
            prs.save(src)
            ed = (settings.get("hymn_edits") or {}).get(b["path"])
            if ed and ed.get("texts"):
                ed["texts"] = {str(int(k) - (int(k) > i)): v for k, v in ed["texts"].items() if int(k) != i}
                save_settings()
            _hymn_cache.clear()
            return self._json({"ok": True, "edit": ed})
        if p == "/api/hymn/blocks":
            return self._json({"blocks": pptx_render.read_blocks(R(self._body()["src"]))})
        if p == "/api/hymn/export":
            # uno o tutti gli inni in una cartella scelta; quelli modificati
            # nell'app vengono rigenerati con le modifiche
            b = self._body()
            dest = pick("folder", "Dove salvo i PowerPoint?")
            if not dest:
                return self._json({"ok": False})
            edits = settings.get("hymn_edits") or {}
            n, nfonts = 0, 0
            have = {}   # famiglia -> {stile: file}, dai caratteri caricati nell'app
            for f in fonts_loaded():
                st = f["style"].lower()
                have.setdefault(f["family"].lower(), {"name": f["family"]})[("boldItalic" if "italic" in st else "bold") if "bold" in st else "italic" if "italic" in st else "regular"] = R(f["path"])
            for h in hymns_for(b["db"]):
                if b.get("only") and h["path"] != b["only"]:
                    continue
                src, out = R(h["path"]), os.path.join(dest[0], os.path.basename(R(h["path"])))
                e = edits.get(h["path"])
                if e and src.lower().endswith(".pptx"):
                    pptx_render.apply_edits(src, out, e)
                else:
                    shutil.copy2(src, out)
                fams = {r["f"] for sl in pptx_render.describe(out, SLIDES)["slides"] for sh in sl if "text" in sh
                        for pa in sh["text"]["paras"] for r in pa["runs"] if r["f"]}
                emb = {have[f.lower()]["name"]: {k: v for k, v in have[f.lower()].items() if k != "name"} for f in fams if f.lower() in have}
                try:
                    nfonts += bool(emb and pptx_render.embed_fonts(out, emb))
                except Exception:
                    pass
                n += 1
            return self._json({"ok": True, "count": n, "folder": dest[0], "fonts": nfonts})
        if p == "/api/update/install":
            u = check_update()
            if not u.get("newer"):
                return self._json({"ok": False, "error": "Hai già l'ultima versione"})
            try:
                install_update(u["asset"])
            except Exception as ex:
                return self._json({"ok": False, "error": str(ex)})
            return self._json({"ok": True})
        if p == "/api/backup/export":
            dest = pick("folder", "Dove salvo il backup?")
            return self._json({"ok": True, "path": backup_to(dest[0])} if dest else {"ok": False})
        if p == "/api/backup/import":
            src = pick("file", "Scegli il backup (.zip)", ["public.zip-archive"])
            if not src:
                return self._json({"ok": False})
            try:
                restore_from(src[0])
            except Exception as ex:
                return self._json({"ok": False, "error": str(ex)})
            return self._json({"ok": True})
        if p == "/api/migrate":
            src = pick("folder", "Scegli la cartella «config» della vecchia app (contiene settings.json)")
            if not src:
                return self._json({"ok": False})
            if not os.path.isfile(os.path.join(src[0], "settings.json")):
                return self._json({"ok": False, "error": "In questa cartella non c'è settings.json"})
            return self._json(migrate_from(src[0]))
        if p == "/api/library/rename":
            b = self._body()
            src = R(b.get("path", ""))
            name = os.path.basename(str(b.get("name", "")).strip()).replace("/", "-")
            if not name or not src.startswith(os.path.join(DATA, "library") + os.sep) or not os.path.isfile(src):
                return self._json({"ok": False}, 400)
            dest = _unique(os.path.join(os.path.dirname(src), name + os.path.splitext(src)[1]))
            os.rename(src, dest)
            return self._json({"ok": True, "path": rel(dest)})
        if p == "/api/hymn/setbase":
            # un audio diventa la base di un inno: copiato nella cartella
            # dell'innario con il numero davanti; la base di prima viene tolta
            b = self._body()
            folder, src, num = R(b.get("db", "")), R(b.get("src", "")), str(b.get("number", ""))
            if not os.path.isdir(folder) or not os.path.isfile(src) or not num.isdigit():
                return self._json({"ok": False}, 400)
            old = audio_by_number(folder).get(int(num))
            while old:
                os.remove(old)
                old = audio_by_number(folder).get(int(num))
            title = str(b.get("title", "")).replace("/", "-")
            dest = os.path.join(folder, f"{num} - {title}{os.path.splitext(src)[1].lower()}")
            shutil.copy2(src, dest)
            _hymn_cache.pop(folder, None)
            return self._json({"ok": True, "audio": rel(dest)})
        if p == "/api/library/import":
            b = self._body()
            if b.get("kind") not in ("video", "audio") or not os.path.isfile(b.get("path", "")):
                return self._json({"ok": False}, 400)
            return self._json({"ok": True, "path": import_into(b["kind"], b["path"])})
        if p == "/api/library/remove":
            target = R(self._body().get("path", ""))
            if target.startswith(os.path.join(DATA, "library") + os.sep) and os.path.exists(target):
                shutil.rmtree(target) if os.path.isdir(target) else os.remove(target)
                _hymn_cache.pop(target, None)
            return self._json({"ok": True})
        self.send_error(404)


def main():
    # tutte le interfacce, così il telefono sulla stessa rete Wi-Fi lo trova;
    # senza PIN impostato il telefono non entra (vedi _gate)
    srv = ThreadingHTTPServer(("0.0.0.0", PORT), Handler)
    srv.daemon_threads = True
    threading.Thread(target=lambda: [hymns_for(d) for d in settings.get("databases", []) if os.path.isdir(R(d))], daemon=True).start()
    global BRIDGE
    try:                                   # telefono con i dati: il ponte gira qui, non nel pannello
        import cloud_bridge
        BRIDGE = cloud_bridge.Bridge(*phone_ids(), PORT).start()
    except Exception as ex:
        print("ponte telefono non avviato:", ex)
    print(f"SDARM {VERSION} — http://127.0.0.1:{PORT} — dati: {DATA}")
    srv.serve_forever()


if __name__ == "__main__":
    main()
